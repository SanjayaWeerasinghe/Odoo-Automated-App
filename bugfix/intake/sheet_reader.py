"""Read bug reports from the shared .xlsx (public export link, read-only)."""
import hashlib
import io
from dataclasses import dataclass, field

import openpyxl
import requests

SKIP_TABS = {"BugSheet - Sample", "Original Sheet"}
# Status -> what the agent does with the row.
ACTIONS = {
    "": "new",
    "need correction": "retry",       # tester says still broken
    "correction done": "recheck",     # dev says fixed; verify
    "completed - new": "recheck",
    "completed": "skip",
}
PRIORITY_ORDER = {"high": 0, "normal": 1, "low": 2}
TEXT_FIELDS = ("transaction", "object_action", "behaviour", "expected", "test_data", "comments")
COLS = {  # canonical field -> header prefix (lowercased, newlines collapsed)
    "date": "date", "reporter": "reported by", "env": "env", "company": "comp", "module": "module",
    "transaction": "transaction", "object_action": "object & action", "behaviour": "behaviour",
    "expected": "expected", "test_data": "test data", "comments": "comments", "priority": "priority",
    "status": "status", "retest": "retest status",
}


@dataclass
class Bug:
    tab: str
    row: int
    fields: dict = field(default_factory=dict)

    @property
    def key(self):
        return f"{self.tab}!{self.row}"

    @property
    def fingerprint(self):
        """Changes when a human edits the report, so edited rows get re-triaged."""
        blob = "|".join(str(self.fields.get(f) or "") for f in TEXT_FIELDS)
        return hashlib.sha1(blob.encode()).hexdigest()[:12]

    @property
    def action(self):
        return ACTIONS.get((self.fields.get("status") or "").strip().lower(), "unknown-status")

    @property
    def priority_rank(self):
        return PRIORITY_ORDER.get((self.fields.get("priority") or "normal").strip().lower(), 1)


def export_url(sheet_id: str) -> str:
    return f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"


def download(sheet_id: str, timeout=120) -> bytes:
    r = requests.get(export_url(sheet_id), timeout=timeout)
    r.raise_for_status()
    return r.content


def parse(xlsx_bytes: bytes) -> list[Bug]:
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    bugs = []
    for ws in wb.worksheets:
        if ws.title.startswith("#") or ws.title in SKIP_TABS or ws.sheet_state != "visible":
            continue
        hdr = [" ".join(str(c.value or "").split()).lower() for c in ws[1]]
        idx = {}
        for name, prefix in COLS.items():
            i = next((i for i, h in enumerate(hdr) if h.startswith(prefix)), None)
            if i is not None:
                idx[name] = i
        for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            f = {k: (str(row[i]).strip() if row[i] not in (None, "") else None) for k, i in idx.items()}
            if not any(f.get(k) for k in TEXT_FIELDS if k != "comments"):
                continue  # empty template row
            bugs.append(Bug(ws.title, r, f))
    return bugs


def work_queue(bugs: list[Bug]) -> list[Bug]:
    """Rows the agent should act on, high priority first, oldest first."""
    todo = [b for b in bugs if b.action in ("new", "retry", "recheck")]
    return sorted(todo, key=lambda b: (b.priority_rank, b.fields.get("date") or "", b.tab, b.row))
